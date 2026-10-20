from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


BASE_DIR = Path(__file__).resolve().parent
OUTPUT = BASE_DIR / 'sgdp_debarquements.xlsx'
MAX_ROWS = 501

GREEN = '087F5B'
DARK_GREEN = '075B45'
YELLOW = 'F2C14E'
LIGHT_GREEN = 'E8F5EF'
LIGHT_YELLOW = 'FFF8DE'
LIGHT_RED = 'FDECEC'
DARK = '173B3F'
WHITE = 'FFFFFF'
GREY = 'F3F5F4'

thin_border = Border(
    left=Side(style='thin', color='D7E2DE'),
    right=Side(style='thin', color='D7E2DE'),
    top=Side(style='thin', color='D7E2DE'),
    bottom=Side(style='thin', color='D7E2DE'),
)


def style_title(sheet, title, subtitle, last_column):
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    sheet.cell(1, 1, title)
    sheet.cell(1, 1).font = Font(size=18, bold=True, color=WHITE)
    sheet.cell(1, 1).fill = PatternFill('solid', fgColor=DARK_GREEN)
    sheet.cell(1, 1).alignment = Alignment(vertical='center')
    sheet.row_dimensions[1].height = 32
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    sheet.cell(2, 1, subtitle)
    sheet.cell(2, 1).font = Font(italic=True, color='5B6D68')
    sheet.row_dimensions[2].height = 24


def style_headers(sheet, row, columns):
    for column in range(1, columns + 1):
        cell = sheet.cell(row, column)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill('solid', fgColor=GREEN)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border
    sheet.row_dimensions[row].height = 30


def add_table(sheet, name, ref):
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name='TableStyleMedium4',
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def create_workbook():
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = 'Instructions'
    landings = workbook.create_sheet('Débarquements')
    sites = workbook.create_sheet('Sites')
    species = workbook.create_sheet('Espèces')
    dashboard = workbook.create_sheet('Tableau de bord')
    lists = workbook.create_sheet('_Listes')
    lists.sheet_state = 'hidden'

    style_title(instructions, 'SGDP - Registre Excel', 'Fichier autonome de saisie des débarquements', 5)
    instructions.column_dimensions['A'].width = 24
    instructions.column_dimensions['B'].width = 92
    instructions['A4'] = 'Démarrage'
    instructions['A4'].font = Font(bold=True, color=GREEN, size=13)
    steps = [
        ('1', 'Ouvrir la feuille Sites et remplacer les exemples par les sites de votre service.'),
        ('2', 'Ouvrir la feuille Espèces et renseigner les espèces autorisées.'),
        ('3', 'Saisir chaque débarquement dans la feuille Débarquements.'),
        ('4', 'Utiliser les listes déroulantes pour le site, l’espèce et le statut.'),
        ('5', 'Consulter le Tableau de bord pour les totaux et les graphiques.'),
    ]
    for row, (number, text) in enumerate(steps, start=5):
        instructions.cell(row, 1, number).font = Font(bold=True, color=GREEN)
        instructions.cell(row, 2, text)
        instructions.cell(row, 2).alignment = Alignment(wrap_text=True, vertical='top')
        instructions.row_dimensions[row].height = 30
    instructions['A12'] = 'Règles de saisie'
    instructions['A12'].font = Font(bold=True, color=GREEN, size=13)
    rules = [
        'Ne pas supprimer les colonnes de la feuille Débarquements.',
        'La quantité doit être exprimée en kilogrammes et supérieure à zéro.',
        'Enregistrer régulièrement le fichier sur l’ordinateur du service.',
        'Ce fichier fonctionne séparément du site Django et ne synchronise pas automatiquement les données.',
    ]
    for row, text in enumerate(rules, start=13):
        instructions.cell(row, 1, '•').font = Font(color=GREEN)
        instructions.cell(row, 2, text)
        instructions.cell(row, 2).alignment = Alignment(wrap_text=True, vertical='top')
        instructions.row_dimensions[row].height = 28

    style_title(landings, 'Débarquements', 'Saisir un débarquement par ligne', 7)
    headers = ['Date', 'Site', 'Espèce', 'Quantité (kg)', 'Observation', 'Statut', 'Saisi par']
    for column, header in enumerate(headers, start=1):
        landings.cell(4, column, header)
    style_headers(landings, 4, len(headers))
    widths = [15, 25, 25, 18, 42, 16, 24]
    for column, width in enumerate(widths, start=1):
        landings.column_dimensions[chr(64 + column)].width = width
    for row in range(5, MAX_ROWS + 1):
        landings.cell(row, 6, 'En attente')
        landings.cell(row, 1).number_format = 'dd/mm/yyyy'
        landings.cell(row, 4).number_format = '#,##0.00'
        for column in range(1, 8):
            landings.cell(row, column).border = thin_border
            landings.cell(row, column).alignment = Alignment(vertical='top', wrap_text=column in (5, 7))
    landings.freeze_panes = 'A5'
    landings.auto_filter.ref = f'A4:G{MAX_ROWS}'
    add_table(landings, 'TableDebarquements', f'A4:G{MAX_ROWS}')

    date_validation = DataValidation(type='date', operator='between', formula1='DATE(2000,1,1)', formula2='DATE(2100,12,31)', allow_blank=True)
    quantity_validation = DataValidation(type='decimal', operator='greaterThan', formula1='0', allow_blank=True)
    status_validation = DataValidation(type='list', formula1='=StatutsListe', allow_blank=False)
    landings.add_data_validation(date_validation)
    landings.add_data_validation(quantity_validation)
    landings.add_data_validation(status_validation)
    date_validation.add(f'A5:A{MAX_ROWS}')
    quantity_validation.add(f'D5:D{MAX_ROWS}')
    status_validation.add(f'F5:F{MAX_ROWS}')

    style_title(sites, 'Sites', 'Référentiel des sites de débarquement', 3)
    for column, header in enumerate(['Nom du site', 'Commune', 'Région'], start=1):
        sites.cell(4, column, header)
    style_headers(sites, 4, 3)
    for column, width in zip('ABC', [32, 28, 28]):
        sites.column_dimensions[column].width = width
    sites['A5'] = 'Exemple à remplacer'
    sites['B5'] = 'Commune'
    sites['C5'] = 'Région'
    sites.freeze_panes = 'A5'
    add_table(sites, 'TableSites', 'A4:C104')
    for row in range(5, 105):
        for column in range(1, 4):
            sites.cell(row, column).border = thin_border

    style_title(species, 'Espèces', 'Référentiel des espèces débarquées', 2)
    for column, header in enumerate(['Nom de l’espèce', 'Catégorie'], start=1):
        species.cell(4, column, header)
    style_headers(species, 4, 2)
    species.column_dimensions['A'].width = 34
    species.column_dimensions['B'].width = 32
    species['A5'] = 'Exemple à remplacer'
    species['B5'] = 'Catégorie'
    species.freeze_panes = 'A5'
    add_table(species, 'TableEspeces', 'A4:B104')
    for row in range(5, 105):
        for column in range(1, 3):
            species.cell(row, column).border = thin_border

    lists['A1'] = 'En attente'
    lists['A2'] = 'Validé'
    workbook.defined_names.add(DefinedName('SitesListe', attr_text="'Sites'!$A$5:$A$104"))
    workbook.defined_names.add(DefinedName('EspecesListe', attr_text="'Espèces'!$A$5:$A$104"))
    workbook.defined_names.add(DefinedName('StatutsListe', attr_text="'_Listes'!$A$1:$A$2"))
    site_validation = DataValidation(type='list', formula1='=SitesListe', allow_blank=False)
    species_validation = DataValidation(type='list', formula1='=EspecesListe', allow_blank=False)
    landings.add_data_validation(site_validation)
    landings.add_data_validation(species_validation)
    site_validation.add(f'B5:B{MAX_ROWS}')
    species_validation.add(f'C5:C{MAX_ROWS}')

    style_title(dashboard, 'Tableau de bord', 'Les indicateurs se mettent à jour depuis la feuille Débarquements', 8)
    dashboard.column_dimensions['A'].width = 30
    dashboard.column_dimensions['B'].width = 20
    dashboard.column_dimensions['D'].width = 30
    dashboard.column_dimensions['E'].width = 20
    dashboard['A4'] = 'Indicateurs'
    dashboard['A4'].font = Font(bold=True, color=GREEN, size=13)
    metrics = [
        ('Nombre de débarquements', f'=COUNTA(\'Débarquements\'!A5:A{MAX_ROWS})'),
        ('Quantité totale (kg)', f'=SUM(\'Débarquements\'!D5:D{MAX_ROWS})'),
        ('Quantité moyenne (kg)', f'=IFERROR(AVERAGE(\'Débarquements\'!D5:D{MAX_ROWS}),0)'),
        ('Débarquements validés', f'=COUNTIF(\'Débarquements\'!F5:F{MAX_ROWS},"Validé")'),
        ('Débarquements en attente', f'=COUNTIF(\'Débarquements\'!F5:F{MAX_ROWS},"En attente")'),
    ]
    for row, (label, formula) in enumerate(metrics, start=5):
        dashboard.cell(row, 1, label)
        dashboard.cell(row, 2, formula)
        dashboard.cell(row, 1).font = Font(bold=True, color=DARK)
        dashboard.cell(row, 2).font = Font(bold=True, color=GREEN, size=12)
        dashboard.cell(row, 1).fill = PatternFill('solid', fgColor=LIGHT_GREEN)
        dashboard.cell(row, 2).fill = PatternFill('solid', fgColor=LIGHT_YELLOW)
        dashboard.cell(row, 1).border = thin_border
        dashboard.cell(row, 2).border = thin_border
    dashboard['B6'].number_format = '#,##0.00'
    dashboard['B7'].number_format = '#,##0.00'

    dashboard['D4'] = 'Quantité par espèce'
    dashboard['D4'].font = Font(bold=True, color=GREEN, size=13)
    dashboard['D5'] = 'Espèce'
    dashboard['E5'] = 'Total (kg)'
    style_headers(dashboard, 5, 5)
    for row in range(6, 26):
        source_row = row - 1
        dashboard.cell(row, 4, f'=IF(\'Espèces\'!A{source_row}="","",\'Espèces\'!A{source_row})')
        dashboard.cell(row, 5, f'=IF(D{row}="","",SUMIF(\'Débarquements\'!C:C,D{row},\'Débarquements\'!D:D))')
        dashboard.cell(row, 5).number_format = '#,##0.00'
    dashboard['G4'] = 'Quantité par site'
    dashboard['G4'].font = Font(bold=True, color=GREEN, size=13)
    dashboard['G5'] = 'Site'
    dashboard['H5'] = 'Total (kg)'
    for cell in dashboard[5][6:8]:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill('solid', fgColor=GREEN)
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border
    for row in range(6, 26):
        source_row = row - 1
        dashboard.cell(row, 7, f'=IF(\'Sites\'!A{source_row}="","",\'Sites\'!A{source_row})')
        dashboard.cell(row, 8, f'=IF(G{row}="","",SUMIF(\'Débarquements\'!B:B,G{row},\'Débarquements\'!D:D))')
        dashboard.cell(row, 8).number_format = '#,##0.00'
    dashboard.freeze_panes = 'A5'

    chart = BarChart()
    chart.title = 'Quantité par espèce'
    chart.y_axis.title = 'Kg'
    chart.x_axis.title = 'Espèce'
    chart.add_data(Reference(dashboard, min_col=5, min_row=5, max_row=25), titles_from_data=True)
    chart.set_categories(Reference(dashboard, min_col=4, min_row=6, max_row=25))
    chart.height = 8
    chart.width = 15
    dashboard.add_chart(chart, 'A13')

    landings.conditional_formatting.add(
        f'F5:F{MAX_ROWS}',
        FormulaRule(formula=['F5="En attente"'], fill=PatternFill('solid', fgColor=LIGHT_YELLOW)),
    )
    landings.conditional_formatting.add(
        f'F5:F{MAX_ROWS}',
        FormulaRule(formula=['F5="Validé"'], fill=PatternFill('solid', fgColor=LIGHT_GREEN)),
    )
    landings.conditional_formatting.add(
        f'D5:D{MAX_ROWS}',
        FormulaRule(formula=['D5<=0'], fill=PatternFill('solid', fgColor=LIGHT_RED)),
    )

    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0

    workbook.active = 1
    workbook.save(OUTPUT)
    print(f'Fichier créé : {OUTPUT}')


if __name__ == '__main__':
    create_workbook()
