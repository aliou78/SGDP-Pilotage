from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName


BASE_DIR = Path(__file__).resolve().parent
OUTPUT = BASE_DIR / 'sgdp_complet.xlsx'
MAX_ROWS = 1004
REF_ROWS = 104

GREEN = '087F5B'
DARK_GREEN = '075B45'
YELLOW = 'F2C14E'
LIGHT_GREEN = 'E8F5EF'
LIGHT_YELLOW = 'FFF8DE'
LIGHT_RED = 'FDECEC'
DARK = '173B3F'
WHITE = 'FFFFFF'
GREY = 'F3F5F4'
BORDER = Border(
    left=Side(style='thin', color='D7E2DE'),
    right=Side(style='thin', color='D7E2DE'),
    top=Side(style='thin', color='D7E2DE'),
    bottom=Side(style='thin', color='D7E2DE'),
)


def title(sheet, name, subtitle, columns):
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    cell = sheet.cell(1, 1, name)
    cell.font = Font(size=18, bold=True, color=WHITE)
    cell.fill = PatternFill('solid', fgColor=DARK_GREEN)
    cell.alignment = Alignment(vertical='center')
    sheet.row_dimensions[1].height = 32
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    cell = sheet.cell(2, 1, subtitle)
    cell.font = Font(italic=True, color='5B6D68')
    sheet.row_dimensions[2].height = 24


def headers(sheet, row, values):
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row, column, value)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill('solid', fgColor=GREEN)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER
    sheet.row_dimensions[row].height = 30


def table(sheet, name, reference):
    item = Table(displayName=name, ref=reference)
    item.tableStyleInfo = TableStyleInfo(
        name='TableStyleMedium4',
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(item)


def validation(sheet, kind, formula, reference):
    item = DataValidation(type=kind, formula1=formula, allow_blank=True)
    sheet.add_data_validation(item)
    item.add(reference)


def format_data(sheet, first_row, last_row, columns):
    for row in range(first_row, last_row + 1):
        for column in range(1, columns + 1):
            cell = sheet.cell(row, column)
            cell.border = BORDER
            cell.alignment = Alignment(vertical='top', wrap_text=column in (6, 8, 9))


def create_workbook():
    book = Workbook()
    accueil = book.active
    accueil.title = 'Accueil'
    connection = book.create_sheet('Connexion', 0)
    users = book.create_sheet('Utilisateurs')
    user_form = book.create_sheet('Formulaire Utilisateurs')
    sites = book.create_sheet('Sites')
    species = book.create_sheet('Espèces')
    site_form = book.create_sheet('Formulaire Sites')
    species_form = book.create_sheet('Formulaire Espèces')
    landings = book.create_sheet('Débarquements')
    reports = book.create_sheet('Rapports')
    dashboard = book.create_sheet('Tableau de bord')
    lists = book.create_sheet('_Listes')
    lists.sheet_state = 'hidden'

    title(connection, 'SGDP - Connexion', 'Page de connexion au fichier Excel', 6)
    connection.column_dimensions['A'].width = 30
    connection.column_dimensions['B'].width = 28
    connection['A5'] = 'Nom utilisateur'
    connection['A6'] = 'Mot de passe'
    connection['B5'] = 'admin'
    connection['B6'] = 'admin123'
    connection['A5'].font = Font(bold=True, color=DARK)
    connection['A6'].font = Font(bold=True, color=DARK)
    connection['A5'].fill = PatternFill('solid', fgColor=LIGHT_GREEN)
    connection['A6'].fill = PatternFill('solid', fgColor=LIGHT_GREEN)
    connection['B5'].fill = PatternFill('solid', fgColor=LIGHT_YELLOW)
    connection['B6'].fill = PatternFill('solid', fgColor=LIGHT_YELLOW)
    connection['A5'].border = BORDER
    connection['A6'].border = BORDER
    connection['B5'].border = BORDER
    connection['B6'].border = BORDER
    connection['A7'] = 'Compte de démonstration : admin / admin123'
    connection['A8'] = 'Dans la version XLSM, utilisez le bouton de connexion après activation des macros.'
    connection['A7'].font = Font(italic=True, color='5B6D68')
    connection['A8'].font = Font(italic=True, color='5B6D68')

    title(accueil, 'SGDP - Gestion Excel complète', 'Version autonome pour la saisie et le suivi des débarquements', 6)
    accueil.column_dimensions['A'].width = 25
    accueil.column_dimensions['B'].width = 95
    accueil['A4'] = 'Fonctionnalités'
    accueil['A4'].font = Font(bold=True, color=GREEN, size=13)
    features = [
        'Utilisateurs : liste des agents, chefs, administrateurs et comptes de consultation.',
        'Sites et espèces : référentiels modifiables utilisés dans les listes déroulantes.',
        'Débarquements : saisie, agent responsable, statut en attente ou validé, validation par un chef.',
        'Rapports : filtres par date, mois, année, site et espèce, avec totaux calculés automatiquement.',
        'Tableau de bord : indicateurs, quantités par site et par espèce, et graphique.',
    ]
    for row, text in enumerate(features, start=5):
        accueil.cell(row, 1, '•').font = Font(color=GREEN)
        accueil.cell(row, 2, text)
        accueil.cell(row, 2).alignment = Alignment(wrap_text=True, vertical='top')
        accueil.row_dimensions[row].height = 30
    accueil['A12'] = 'Mode d’utilisation'
    accueil['A12'].font = Font(bold=True, color=GREEN, size=13)
    notes = [
        'Renseigner d’abord les feuilles Sites et Espèces.',
        'Saisir ensuite les débarquements dans la feuille Débarquements.',
        'Pour valider un débarquement, remplacer son statut par Validé et renseigner Valide par et Date validation.',
        'Les feuilles Rapports et Tableau de bord se recalculent à l’ouverture dans Excel.',
        'Ce fichier est autonome et ne se synchronise pas avec le site Django.',
        'Excel ne fournit pas la même sécurité de connexion et de permissions que le site web.',
    ]
    for row, text in enumerate(notes, start=13):
        accueil.cell(row, 1, '•').font = Font(color=GREEN)
        accueil.cell(row, 2, text)
        accueil.cell(row, 2).alignment = Alignment(wrap_text=True, vertical='top')
        accueil.row_dimensions[row].height = 28

    title(users, 'Utilisateurs', 'Référentiel des personnes autorisées à utiliser le fichier', 8)
    user_headers = ['Nom utilisateur', 'Prénom', 'Nom', 'Rôle', 'Téléphone', 'Adresse', 'Compte actif', 'Mot de passe']
    headers(users, 4, user_headers)
    for column, width in enumerate([24, 20, 20, 22, 18, 34, 16, 22], start=1):
        users.column_dimensions[chr(64 + column)].width = width
    format_data(users, 5, REF_ROWS, len(user_headers))
    users.freeze_panes = 'A5'
    table(users, 'TableUtilisateurs', f'A4:H{REF_ROWS}')
    validation(users, 'list', '=RolesListe', f'D5:D{REF_ROWS}')
    validation(users, 'list', '=ActifListe', f'G5:G{REF_ROWS}')

    title(sites, 'Sites', 'Référentiel des sites de débarquement', 4)
    site_headers = ['Nom du site', 'Commune', 'Région', 'Actif']
    headers(sites, 4, site_headers)
    for column, width in enumerate([32, 28, 28, 14], start=1):
        sites.column_dimensions[chr(64 + column)].width = width
    format_data(sites, 5, REF_ROWS, len(site_headers))
    sites.freeze_panes = 'A5'
    table(sites, 'TableSites', f'A4:D{REF_ROWS}')
    validation(sites, 'list', '=ActifListe', f'D5:D{REF_ROWS}')

    title(species, 'Espèces', 'Référentiel des espèces débarquées', 3)
    species_headers = ['Nom de l’espèce', 'Nom scientifique', 'Catégorie', 'Actif']
    headers(species, 4, species_headers)
    for column, width in enumerate([34, 34, 32, 14], start=1):
        species.column_dimensions[chr(64 + column)].width = width
    format_data(species, 5, REF_ROWS, len(species_headers))
    species.freeze_panes = 'A5'
    table(species, 'TableEspeces', f'A4:D{REF_ROWS}')
    validation(species, 'list', '=ActifListe', f'D5:D{REF_ROWS}')

    title(user_form, 'Formulaire Utilisateurs', 'Créer, rechercher et gérer les comptes Excel', 4)
    user_form.column_dimensions['A'].width = 28
    user_form.column_dimensions['B'].width = 36
    user_form.column_dimensions['C'].width = 20
    user_form['A4'] = 'Nom utilisateur'
    user_form['A5'] = 'Prénom'
    user_form['A6'] = 'Nom'
    user_form['A7'] = 'Rôle'
    user_form['A8'] = 'Téléphone'
    user_form['A9'] = 'Adresse'
    user_form['A10'] = 'Compte actif'
    user_form['A11'] = 'Mot de passe'
    for row in range(4, 12):
        user_form.cell(row, 1).font = Font(bold=True, color=DARK)
        user_form.cell(row, 1).fill = PatternFill('solid', fgColor=LIGHT_GREEN)
        user_form.cell(row, 1).border = BORDER
        user_form.cell(row, 2).fill = PatternFill('solid', fgColor=LIGHT_YELLOW)
        user_form.cell(row, 2).border = BORDER
    user_form['A13'] = 'Le bouton de recherche utilise le nom utilisateur, le prénom, le nom ou le rôle.'
    user_form['A13'].font = Font(italic=True, color='5B6D68')

    title(site_form, 'Recherche de sites', 'Rechercher un site dans le référentiel', 4)
    site_form.column_dimensions['A'].width = 28
    site_form.column_dimensions['B'].width = 34
    site_form['A4'] = 'Recherche'
    for row in range(4, 5):
        site_form.cell(row, 1).font = Font(bold=True, color=DARK)
        site_form.cell(row, 1).fill = PatternFill('solid', fgColor=LIGHT_GREEN)
        site_form.cell(row, 1).border = BORDER
        site_form.cell(row, 2).fill = PatternFill('solid', fgColor=LIGHT_YELLOW)
        site_form.cell(row, 2).border = BORDER
    site_form['A6'] = 'Laissez vide pour afficher tous les sites.'
    site_form['A6'].font = Font(italic=True, color='5B6D68')
    site_form['A8'] = 'La recherche porte sur le nom, la commune et la région.'
    site_form['A8'].font = Font(italic=True, color='5B6D68')

    title(species_form, 'Recherche d’espèces', 'Rechercher une espèce dans le référentiel', 4)
    species_form.column_dimensions['A'].width = 28
    species_form.column_dimensions['B'].width = 34
    species_form['A4'] = 'Recherche'
    for row in range(4, 5):
        species_form.cell(row, 1).font = Font(bold=True, color=DARK)
        species_form.cell(row, 1).fill = PatternFill('solid', fgColor=LIGHT_GREEN)
        species_form.cell(row, 1).border = BORDER
        species_form.cell(row, 2).fill = PatternFill('solid', fgColor=LIGHT_YELLOW)
        species_form.cell(row, 2).border = BORDER
    species_form['A6'] = 'Laissez vide pour afficher toutes les espèces.'
    species_form['A6'].font = Font(italic=True, color='5B6D68')
    species_form['A7'] = 'La recherche porte sur le nom et la catégorie.'
    species_form['A7'].font = Font(italic=True, color='5B6D68')

    title(landings, 'Débarquements', 'Saisie, suivi et validation des débarquements', 10)
    landing_headers = [
        'ID', 'Date', 'Site', 'Espèce', 'Quantité (kg)', 'Observation',
        'Statut', 'Agent', 'Valide par', 'Date validation',
    ]
    headers(landings, 4, landing_headers)
    for column, width in enumerate([10, 15, 26, 28, 18, 42, 16, 24, 24, 20], start=1):
        landings.column_dimensions[chr(64 + column)].width = width
    for row in range(5, MAX_ROWS + 1):
        landings.cell(row, 1, row - 4)
        landings.cell(row, 7, 'En attente')
        landings.cell(row, 2).number_format = 'dd/mm/yyyy'
        landings.cell(row, 5).number_format = '#,##0.00'
        landings.cell(row, 10).number_format = 'dd/mm/yyyy'
    format_data(landings, 5, MAX_ROWS, len(landing_headers))
    landings.freeze_panes = 'A5'
    table(landings, 'TableDebarquements', f'A4:J{MAX_ROWS}')
    validation(landings, 'date', 'DATE(2000,1,1)', f'B5:B{MAX_ROWS}')
    validation(landings, 'list', '=SitesListe', f'C5:C{MAX_ROWS}')
    validation(landings, 'list', '=EspecesListe', f'D5:D{MAX_ROWS}')
    validation(landings, 'decimal', '0', f'E5:E{MAX_ROWS}')
    validation(landings, 'list', '=StatutsListe', f'G5:G{MAX_ROWS}')
    validation(landings, 'list', '=UtilisateursListe', f'H5:H{MAX_ROWS}')
    validation(landings, 'list', '=UtilisateursListe', f'I5:I{MAX_ROWS}')
    validation(landings, 'date', 'DATE(2000,1,1)', f'J5:J{MAX_ROWS}')
    landings.conditional_formatting.add(
        f'G5:G{MAX_ROWS}',
        FormulaRule(formula=['G5="En attente"'], fill=PatternFill('solid', fgColor=LIGHT_YELLOW)),
    )
    landings.conditional_formatting.add(
        f'G5:G{MAX_ROWS}',
        FormulaRule(formula=['G5="Validé"'], fill=PatternFill('solid', fgColor=LIGHT_GREEN)),
    )
    landings.conditional_formatting.add(
        f'E5:E{MAX_ROWS}',
        FormulaRule(formula=['E5<=0'], fill=PatternFill('solid', fgColor=LIGHT_RED)),
    )

    lists['A1'] = 'En attente'
    lists['A2'] = 'Validé'
    lists['B1'] = 'Actif'
    lists['B2'] = 'Inactif'
    lists['C1'] = 'admin'
    lists['C2'] = 'chef'
    lists['C3'] = 'agent'
    lists['C4'] = 'consultation'
    for row in range(1, 101):
        lists.cell(row, 4, f'=IF(Utilisateurs!A{row + 4}="","",Utilisateurs!A{row + 4})')
        lists.cell(row, 5, f'=IF(Sites!A{row + 4}="","",Sites!A{row + 4})')
        lists.cell(row, 6, f'=IF(Espèces!A{row + 4}="","",Espèces!A{row + 4})')
    book.defined_names.add(DefinedName('StatutsListe', attr_text="'_Listes'!$A$1:$A$2"))
    book.defined_names.add(DefinedName('ActifListe', attr_text="'_Listes'!$B$1:$B$2"))
    book.defined_names.add(DefinedName('RolesListe', attr_text="'_Listes'!$C$1:$C$4"))
    book.defined_names.add(DefinedName('UtilisateursListe', attr_text="'_Listes'!$D$1:$D$100"))
    book.defined_names.add(DefinedName('SitesListe', attr_text="'_Listes'!$E$1:$E$100"))
    book.defined_names.add(DefinedName('EspecesListe', attr_text="'_Listes'!$F$1:$F$100"))

    title(reports, 'Rapports', 'Filtres et indicateurs des débarquements', 8)
    reports.column_dimensions['A'].width = 30
    reports.column_dimensions['B'].width = 22
    reports.column_dimensions['D'].width = 30
    reports.column_dimensions['E'].width = 20
    reports['A4'] = 'Filtres du rapport'
    reports['A4'].font = Font(bold=True, color=GREEN, size=13)
    report_filters = [
        ('Type de rapport', 'Journalier'),
        ('Date', ''),
        ('Mois (1-12)', ''),
        ('Année', ''),
        ('Site', ''),
        ('Espèce', ''),
    ]
    for row, (label, value) in enumerate(report_filters, start=5):
        reports.cell(row, 1, label).font = Font(bold=True, color=DARK)
        reports.cell(row, 2, value)
        reports.cell(row, 1).fill = PatternFill('solid', fgColor=LIGHT_GREEN)
        reports.cell(row, 2).fill = PatternFill('solid', fgColor=LIGHT_YELLOW)
        reports.cell(row, 1).border = BORDER
        reports.cell(row, 2).border = BORDER
    validation(reports, 'list', '"Journalier,Mensuel,Annuel,Personnalisé"', 'B5')
    validation(reports, 'date', 'DATE(2000,1,1)', 'B6')
    validation(reports, 'whole', '1', 'B7')
    validation(reports, 'whole', '2000', 'B8')
    validation(reports, 'list', '=SitesListe', 'B9')
    validation(reports, 'list', '=EspecesListe', 'B10')
    reports['D4'] = 'Résultats'
    reports['D4'].font = Font(bold=True, color=GREEN, size=13)
    report_metrics = [
        ('Nombre de débarquements', f'=SUMPRODUCT(--(\'Débarquements\'!$B$5:$B${MAX_ROWS}<>""),IF($B$6="",1,--(\'Débarquements\'!$B$5:$B${MAX_ROWS}=$B$6)),IF($B$7="",1,--(MONTH(\'Débarquements\'!$B$5:$B${MAX_ROWS})=$B$7)),IF($B$8="",1,--(YEAR(\'Débarquements\'!$B$5:$B${MAX_ROWS})=$B$8)),IF($B$9="",1,--(\'Débarquements\'!$C$5:$C${MAX_ROWS}=$B$9)),IF($B$10="",1,--(\'Débarquements\'!$D$5:$D${MAX_ROWS}=$B$10)))'),
        ('Quantité totale (kg)', f'=SUMPRODUCT(\'Débarquements\'!$E$5:$E${MAX_ROWS},IF($B$6="",1,--(\'Débarquements\'!$B$5:$B${MAX_ROWS}=$B$6)),IF($B$7="",1,--(MONTH(\'Débarquements\'!$B$5:$B${MAX_ROWS})=$B$7)),IF($B$8="",1,--(YEAR(\'Débarquements\'!$B$5:$B${MAX_ROWS})=$B$8)),IF($B$9="",1,--(\'Débarquements\'!$C$5:$C${MAX_ROWS}=$B$9)),IF($B$10="",1,--(\'Débarquements\'!$D$5:$D${MAX_ROWS}=$B$10)))'),
        ('Quantité validée (kg)', f'=SUMPRODUCT(\'Débarquements\'!$E$5:$E${MAX_ROWS},--(\'Débarquements\'!$G$5:$G${MAX_ROWS}="Validé"),IF($B$6="",1,--(\'Débarquements\'!$B$5:$B${MAX_ROWS}=$B$6)),IF($B$7="",1,--(MONTH(\'Débarquements\'!$B$5:$B${MAX_ROWS})=$B$7)),IF($B$8="",1,--(YEAR(\'Débarquements\'!$B$5:$B${MAX_ROWS})=$B$8)),IF($B$9="",1,--(\'Débarquements\'!$C$5:$C${MAX_ROWS}=$B$9)),IF($B$10="",1,--(\'Débarquements\'!$D$5:$D${MAX_ROWS}=$B$10)))'),
    ]
    for row, (label, formula) in enumerate(report_metrics, start=5):
        reports.cell(row, 4, label).font = Font(bold=True, color=DARK)
        reports.cell(row, 5, formula).font = Font(bold=True, color=GREEN, size=12)
        reports.cell(row, 4).fill = PatternFill('solid', fgColor=LIGHT_GREEN)
        reports.cell(row, 5).fill = PatternFill('solid', fgColor=LIGHT_YELLOW)
        reports.cell(row, 4).border = BORDER
        reports.cell(row, 5).border = BORDER
    reports['E6'].number_format = '#,##0.00'
    reports['E7'].number_format = '#,##0.00'
    reports['D11'] = 'Détail par espèce'
    reports['D11'].font = Font(bold=True, color=GREEN, size=13)
    reports['D12'] = 'Espèce'
    reports['E12'] = 'Quantité (kg)'
    headers(reports, 12, ['Détail par espèce', '', '', 'Espèce', 'Quantité (kg)'])
    for row in range(13, 33):
        source = row - 8
        reports.cell(row, 4, f'=IF(Espèces!A{source}="","",Espèces!A{source})')
        reports.cell(
            row,
            5,
            f'=IF(D{row}="","",SUMPRODUCT(\'Débarquements\'!$E$5:$E${MAX_ROWS},--(\'Débarquements\'!$D$5:$D${MAX_ROWS}=D{row}),IF($B$6="",1,--(\'Débarquements\'!$B$5:$B${MAX_ROWS}=$B$6)),IF($B$7="",1,--(MONTH(\'Débarquements\'!$B$5:$B${MAX_ROWS})=$B$7)),IF($B$8="",1,--(YEAR(\'Débarquements\'!$B$5:$B${MAX_ROWS})=$B$8)),IF($B$9="",1,--(\'Débarquements\'!$C$5:$C${MAX_ROWS}=$B$9)),IF($B$10="",1,--(\'Débarquements\'!$D$5:$D${MAX_ROWS}=$B$10))))',
        )
        reports.cell(row, 5).number_format = '#,##0.00'

    reports['G11'] = 'Détail par site'
    reports['G11'].font = Font(bold=True, color=GREEN, size=13)
    reports['G12'] = 'Site'
    reports['H12'] = 'Quantité (kg)'
    for cell in reports[12][6:8]:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill('solid', fgColor=GREEN)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = BORDER
    for row in range(13, 33):
        source = row - 8
        reports.cell(row, 7, f'=IF(Sites!A{source}="","",Sites!A{source})')
        reports.cell(
            row,
            8,
            f'=IF(G{row}="","",SUMPRODUCT(\'Débarquements\'!$E$5:$E${MAX_ROWS},--(\'Débarquements\'!$C$5:$C${MAX_ROWS}=G{row}),IF($B$6="",1,--(\'Débarquements\'!$B$5:$B${MAX_ROWS}=$B$6)),IF($B$7="",1,--(MONTH(\'Débarquements\'!$B$5:$B${MAX_ROWS})=$B$7)),IF($B$8="",1,--(YEAR(\'Débarquements\'!$B$5:$B${MAX_ROWS})=$B$8)),IF($B$9="",1,--(\'Débarquements\'!$C$5:$C${MAX_ROWS}=$B$9)),IF($B$10="",1,--(\'Débarquements\'!$D$5:$D${MAX_ROWS}=$B$10))))',
        )
        reports.cell(row, 8).number_format = '#,##0.00'

    reports.column_dimensions['G'].width = 30
    reports.column_dimensions['H'].width = 20

    species_chart = PieChart()
    species_chart.title = 'Répartition filtrée par espèce'
    species_chart.add_data(
        Reference(reports, min_col=5, min_row=12, max_row=32),
        titles_from_data=True,
    )
    species_chart.set_categories(Reference(reports, min_col=4, min_row=13, max_row=32))
    species_chart.height = 8
    species_chart.width = 13
    reports.add_chart(species_chart, 'A35')

    site_chart = BarChart()
    site_chart.title = 'Quantité filtrée par site'
    site_chart.y_axis.title = 'Quantité (kg)'
    site_chart.x_axis.title = 'Sites'
    site_chart.add_data(
        Reference(reports, min_col=8, min_row=12, max_row=32),
        titles_from_data=True,
    )
    site_chart.set_categories(Reference(reports, min_col=7, min_row=13, max_row=32))
    site_chart.height = 8
    site_chart.width = 15
    reports.add_chart(site_chart, 'G35')

    reports.freeze_panes = 'A5'
    reports.print_area = 'A1:N55'

    title(dashboard, 'Tableau de bord', 'Vue synthétique de l’activité', 8)
    dashboard.column_dimensions['A'].width = 32
    dashboard.column_dimensions['B'].width = 20
    dashboard.column_dimensions['D'].width = 30
    dashboard.column_dimensions['E'].width = 20
    dashboard['A4'] = 'Indicateurs généraux'
    dashboard['A4'].font = Font(bold=True, color=GREEN, size=13)
    general_metrics = [
        ('Sites actifs', '=COUNTIF(Sites!D5:D104,"Actif")'),
        ('Espèces actives', '=COUNTIF(Espèces!D5:D104,"Actif")'),
        ('Débarquements', f'=COUNTA(\'Débarquements\'!B5:B{MAX_ROWS})'),
        ('Quantité totale (kg)', f'=SUM(\'Débarquements\'!E5:E{MAX_ROWS})'),
        ('Débarquements validés', f'=COUNTIF(\'Débarquements\'!G5:G{MAX_ROWS},"Validé")'),
        ('Débarquements en attente', f'=COUNTIF(\'Débarquements\'!G5:G{MAX_ROWS},"En attente")'),
    ]
    for row, (label, formula) in enumerate(general_metrics, start=5):
        dashboard.cell(row, 1, label).font = Font(bold=True, color=DARK)
        dashboard.cell(row, 2, formula).font = Font(bold=True, color=GREEN, size=12)
        dashboard.cell(row, 1).fill = PatternFill('solid', fgColor=LIGHT_GREEN)
        dashboard.cell(row, 2).fill = PatternFill('solid', fgColor=LIGHT_YELLOW)
        dashboard.cell(row, 1).border = BORDER
        dashboard.cell(row, 2).border = BORDER
    dashboard['B8'].number_format = '#,##0.00'
    dashboard['D4'] = 'Quantité par espèce'
    dashboard['D4'].font = Font(bold=True, color=GREEN, size=13)
    dashboard['D5'] = 'Espèce'
    dashboard['E5'] = 'Total (kg)'
    for cell in dashboard[5][3:5]:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill('solid', fgColor=GREEN)
        cell.border = BORDER
    for row in range(6, 26):
        source = row - 1
        dashboard.cell(row, 4, f'=IF(Espèces!A{source}="","",Espèces!A{source})')
        dashboard.cell(row, 5, f'=IF(D{row}="","",SUMIF(\'Débarquements\'!D:D,D{row},\'Débarquements\'!E:E))')
        dashboard.cell(row, 5).number_format = '#,##0.00'
    dashboard['G4'] = 'Quantité par site'
    dashboard['G4'].font = Font(bold=True, color=GREEN, size=13)
    dashboard['G5'] = 'Site'
    dashboard['H5'] = 'Total (kg)'
    for cell in dashboard[5][6:8]:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill('solid', fgColor=GREEN)
        cell.border = BORDER
    for row in range(6, 26):
        source = row - 1
        dashboard.cell(row, 7, f'=IF(Sites!A{source}="","",Sites!A{source})')
        dashboard.cell(row, 8, f'=IF(G{row}="","",SUMIF(\'Débarquements\'!C:C,G{row},\'Débarquements\'!E:E))')
        dashboard.cell(row, 8).number_format = '#,##0.00'
    chart = BarChart()
    chart.title = 'Quantité par espèce'
    chart.y_axis.title = 'Kg'
    chart.x_axis.title = 'Espèce'
    chart.add_data(Reference(dashboard, min_col=5, min_row=5, max_row=25), titles_from_data=True)
    chart.set_categories(Reference(dashboard, min_col=4, min_row=6, max_row=25))
    chart.height = 8
    chart.width = 15
    dashboard.add_chart(chart, 'A13')
    dashboard.freeze_panes = 'A5'

    for sheet in book.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0

    book.active = 0
    book.save(OUTPUT)
    print(f'Fichier créé : {OUTPUT}')


if __name__ == '__main__':
    create_workbook()
