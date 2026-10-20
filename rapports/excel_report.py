from datetime import datetime
from io import BytesIO

from django.contrib.staticfiles import finders
from django.db.models import Sum
from django.http import HttpResponse
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins


def generer_rapport_activites_excel(filtres):
    (
        debarquements,
        type_rapport,
        date_rapport,
        mois,
        annee,
        date_debut,
        date_fin,
        site_id,
        espece_id,
    ) = filtres

    mois_noms = {
        '1': 'janvier',
        '2': 'février',
        '3': 'mars',
        '4': 'avril',
        '5': 'mai',
        '6': 'juin',
        '7': 'juillet',
        '8': 'août',
        '9': 'septembre',
        '10': 'octobre',
        '11': 'novembre',
        '12': 'décembre',
    }
    types_rapport = {
        'journalier': 'Rapport journalier',
        'mensuel': 'Rapport d’activités mensuel',
        'annuel': 'Rapport d’activités annuel',
        'personnalise': 'Rapport d’activités',
    }
    periode = types_rapport.get(type_rapport, 'Rapport d’activités')
    if type_rapport == 'journalier' and date_rapport:
        try:
            periode += ' du ' + datetime.strptime(
                date_rapport, '%Y-%m-%d'
            ).strftime('%d/%m/%Y')
        except ValueError:
            periode += ' du ' + date_rapport
    elif type_rapport == 'mensuel':
        if mois and annee:
            periode += f' — {mois_noms.get(str(mois), mois)} {annee}'
        elif annee:
            periode += f' — {annee}'
    elif type_rapport == 'annuel' and annee:
        periode += f' — {annee}'
    elif type_rapport == 'personnalise':
        periode += f' — du {date_debut or "________"} au {date_fin or "________"}'

    from especes.models import Espece
    from sites.models import Site

    site_selectionne = (
        Site.objects.filter(pk=site_id).values_list('nom', flat=True).first()
        if site_id else ''
    )
    espece_selectionnee = (
        Espece.objects.filter(pk=espece_id).values_list('nom', flat=True).first()
        if espece_id else ''
    )
    resume_sites_especes = list(
        debarquements.values('site__nom', 'espece__nom')
        .annotate(total=Sum('quantite'))
        .order_by('site__nom', 'espece__nom')
    )
    total_general = debarquements.aggregate(
        total=Sum('quantite')
    )['total'] or 0

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Rapport d’activités'
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = 'A12'
    worksheet.page_setup.orientation = 'landscape'
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_margins = PageMargins(
        left=0.25, right=0.25, top=0.4, bottom=0.4, header=0.2, footer=0.2
    )
    worksheet.print_title_rows = '1:11'

    thin_gray = Side(style='thin', color='666666')
    cell_border = Border(
        left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray
    )
    section_fill = PatternFill('solid', fgColor='D9E5DF')
    header_fill = PatternFill('solid', fgColor='EDF2EF')
    bold_font = Font(name='Arial', size=9, bold=True)
    body_font = Font(name='Arial', size=9)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)

    for column, width in {
        'A': 12, 'B': 24, 'C': 22, 'D': 20, 'E': 24, 'F': 24,
    }.items():
        worksheet.column_dimensions[column].width = width

    logo_path = finders.find('images/logo .png')
    if logo_path:
        logo = ExcelImage(logo_path)
        logo.width = 90
        logo.height = 65
        worksheet.add_image(logo, 'A1')

    letterhead = [
        ('RÉPUBLIQUE DU SÉNÉGAL', 1, 12, True),
        ('Un Peuple - Un But - Une Foi', 2, 10, False),
        ('MINISTÈRE DES PÊCHES ET DE L’ÉCONOMIE MARITIME', 4, 10, True),
        ('DIRECTION DE LA PÊCHE CONTINENTALE', 5, 10, True),
        ('SERVICE RÉGIONAL DES PÊCHES ET DE LA SURVEILLANCE — SÉDHIOU', 6, 9, False),
        ('SERVICE DÉPARTEMENTAL DES PÊCHES ET DE LA SURVEILLANCE — BOUNKILING', 7, 9, False),
    ]
    for text, row, size, bold in letterhead:
        worksheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        cell = worksheet.cell(row=row, column=2, value=text)
        cell.font = Font(name='Arial', size=size, bold=bold)
        cell.alignment = center
    worksheet.row_dimensions[3].height = 6
    worksheet.row_dimensions[8].height = 8

    worksheet.merge_cells('A9:F9')
    title = worksheet['A9']
    title.value = periode.upper()
    title.font = Font(name='Arial', size=15, bold=True)
    title.alignment = center
    worksheet.merge_cells('A10:F10')
    worksheet['A10'] = 'SUIVI DES ACTIVITÉS ET DE LA PRODUCTION DE LA PÊCHE'
    worksheet['A10'].font = Font(name='Arial', size=10, bold=True)
    worksheet['A10'].alignment = center
    worksheet.merge_cells('A11:F11')
    filter_labels = []
    if site_selectionne:
        filter_labels.append(f'Site : {site_selectionne}')
    if espece_selectionnee:
        filter_labels.append(f'Espèce : {espece_selectionnee}')
    worksheet['A11'] = ' — '.join(filter_labels)
    worksheet['A11'].font = Font(name='Arial', size=9, italic=True)
    worksheet['A11'].alignment = center

    def add_section(title, headers=None, rows=None):
        rows = rows or []
        row_index = worksheet.max_row + 1
        if title:
            row_index += 1
            worksheet.merge_cells(
                start_row=row_index,
                start_column=1,
                end_row=row_index,
                end_column=6,
            )
            heading = worksheet.cell(row=row_index, column=1, value=title)
            heading.font = Font(name='Arial', size=10, bold=True)
            heading.fill = section_fill
            heading.alignment = left
            for cell in worksheet[row_index][:6]:
                cell.fill = section_fill
                cell.border = cell_border
            row_index += 1

        if headers:
            for column, value in enumerate(headers, start=1):
                cell = worksheet.cell(row=row_index, column=column, value=value)
                cell.font = bold_font
                cell.fill = header_fill
                cell.alignment = center
                cell.border = cell_border
            row_index += 1

        for values in rows:
            if len(values) == 1:
                worksheet.merge_cells(
                    start_row=row_index,
                    start_column=1,
                    end_row=row_index,
                    end_column=6,
                )
            for column, value in enumerate(values, start=1):
                cell = worksheet.cell(row=row_index, column=column, value=value)
                cell.font = Font(name='Arial', size=9, italic=True) if len(values) == 1 else body_font
                cell.alignment = left if column > 1 else center
                cell.border = cell_border
            if len(values) > 1:
                for column in range(len(values) + 1, 7):
                    worksheet.cell(row=row_index, column=column).border = cell_border
            worksheet.row_dimensions[row_index].height = 24 if len(values) == 1 else 20
            row_index += 1
        return row_index

    add_section('INTRODUCTION', rows=[[
        'Rapport des activités du Service Départemental des Pêches et de la '
        'Surveillance de Bounkiling pour la période sélectionnée. Les données '
        'de débarquement sont issues du SGDP. Les rubriques non gérées par '
        'l’application restent à compléter avant validation.',
    ]])
    add_section('1. PERSONNEL', [
        'N°', 'Prénom(s) et nom', 'Profession', 'Fonction',
    ], [['1', '', '', ''], ['2', '', '', ''], ['3', '', '', '']])
    add_section('2. LOGISTIQUE', [
        'Désignation', 'État', 'Besoins',
    ], [['', '', ''], ['', '', '']])
    add_section('3. BESOINS MAJEURS', rows=[[''], ['']])
    add_section('4. ÉVÉNEMENT MAJEUR DANS LE MOIS', rows=[['']])
    add_section('5. ACTIVITÉS RÉALISÉES', [
        'N°', 'Activité', 'Observations',
    ], [['1', '', ''], ['2', '', ''], ['3', '', ''], ['4', '', '']])
    add_section('6. SITUATION DES CARTES PROFESSIONNELLES ET PERMIS DE PÊCHE', [
        'Type de permis / carte', 'Catégorie', 'Nombre', 'Total',
    ], [
        ['Permis national', 'A', '', ''],
        ['Permis national', 'B', '', ''],
        ['Permis national', 'C', '', ''],
        ['Permis étranger', 'A', '', ''],
        ['Permis étranger', 'B', '', ''],
        ['Permis étranger', 'C', '', ''],
        ['Cartes de mareyeurs', 'À compléter', '', ''],
        ['TOTAL', '', '', ''],
    ])
    add_section('7. CONTENTIEUX', [
        'Type d’infraction', 'Nombre', 'Amendes (FCFA)',
    ], [
        ['Pêche', '', ''],
        ['Transformation', '', ''],
        ['Mareyage', '', ''],
        ['Transport fluvial', '', ''],
        ['Total', '', ''],
    ])
    add_section('8. BESOIN EN GILETS DE SAUVETAGE', [
        'Besoin', 'Nombre',
    ], [['Gilets de sauvetage', '']])
    add_section('9. SITUATION DU PARC PIROGUIER', [
        'Désignation', 'Sites', 'Nombre',
    ], [
        ['Pirogue de pêche', '', ''],
        ['Pirogue de transport', '', ''],
        ['TOTAL', '', ''],
    ])
    add_section('10. SITUATION DES ORGANISATIONS PROFESSIONNELLES', [
        'N°', 'Dénomination', 'N° RCCM', 'Membres (H/F)',
        'Responsable / contact', 'Adresse',
    ], [['1', '', '', '', '', ''], ['2', '', '', '', '', ''], ['3', '', '', '', '', '']])
    add_section('11. SÉCURITÉ DE LA PÊCHE ET DU TRANSPORT FLUVIAL',
                rows=[[''], [''], ['']])
    add_section('12. PRODUCTION DE LA PÊCHE LOCALE — MISES À TERRE', rows=[[
        'Quantités calculées depuis les débarquements SGDP. Les prix unitaires '
        'et valeurs commerciales sont à compléter.',
    ]])
    production_rows = [
        [
            ligne['site__nom'] or '',
            ligne['espece__nom'] or '',
            ligne['total'] or 0,
            '',
            '',
        ]
        for ligne in resume_sites_especes
    ]
    if not production_rows:
        production_rows.append(['', 'Aucun débarquement enregistré', '', '', ''])
    production_rows.append(['', 'TOTAL GÉNÉRAL', total_general, '', ''])
    add_section('', [
        'Site', 'Espèce', 'Poids mis à terre (kg)',
        'Prix unitaire (FCFA/kg)', 'Valeur commerciale (FCFA)',
    ], production_rows)
    add_section('13. PHYSIONOMIE GÉNÉRALE DU MAREYAGE', rows=[[''], ['']])
    add_section('14. APPORTS EXTÉRIEURS — PRODUITS FRAIS', [
        'Espèce', 'Quantité (kg)', 'Prix unitaire (FCFA)',
        'Nature', 'Valeur estimée (FCFA)',
    ], [['', '', '', 'Frais', ''], ['', '', '', 'Frais', ''], ['TOTAL', '', '', '', '']])
    add_section('15. MAREYAGE DE PRODUITS TRANSFORMÉS', [
        'Nom vernaculaire', 'Quantité (kg)', 'Prix unitaire (FCFA)',
        'Nature', 'Valeur estimée (FCFA)',
    ], [['', '', '', 'Transformé', ''], ['', '', '', 'Transformé', ''], ['TOTAL', '', '', '', '']])
    add_section('16. PROJETS EN COURS D’EXÉCUTION', rows=[['']])
    add_section('17. PERSPECTIVES', rows=[[''], ['']])
    add_section('18. CONTRAINTES LIÉES AU FONCTIONNEMENT DU SERVICE',
                rows=[[''], ['']])
    add_section('', rows=[
        [f'Fait à Bounkiling, le {timezone.localdate().strftime("%d/%m/%Y")}',
         '', '', 'Le Chef de Service'],
        ['', '', '', ''],
        ['', '', '', ''],
    ])

    worksheet.print_area = f'A1:F{worksheet.max_row}'
    worksheet.oddFooter.center.text = 'SGDP — Rapport d’activités'
    worksheet.oddFooter.right.text = 'Page &[Page] sur &[Pages]'

    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type=(
            'application/vnd.openxmlformats-'
            'officedocument.spreadsheetml.sheet'
        ),
    )
    response['Content-Disposition'] = (
        'attachment; filename="rapport_activites.xlsx"'
    )
    return response
