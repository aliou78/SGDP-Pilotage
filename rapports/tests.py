from io import BytesIO

from django.test import RequestFactory, TestCase
from django.db.models import Sum
from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import User
from debarquements.models import Debarquement
from especes.models import Espece
from sites.models import Site

from .views import (
	appliquer_filtres,
	construire_evolution,
	rapport_excel,
	rapport_pdf,
)


class RapportNavigationTests(TestCase):

	def test_menu_affiche_les_rapports_aux_roles_autorises(self):
		for role in ('admin', 'chef', 'consultation'):
			with self.subTest(role=role):
				user = User.objects.create_user(
					username=f'utilisateur_{role}',
					password='password123',
					role=role,
				)
				self.client.force_login(user)

				response = self.client.get(reverse('rapport_list'))

				self.assertEqual(response.status_code, 200)
				self.assertContains(response, 'href="/rapports/"')

	def test_exports_soumettent_les_filtres_selectionnes(self):
		user = User.objects.create_user(
			username='admin_exports',
			password='password123',
			role='admin',
		)
		self.client.force_login(user)

		response = self.client.get(reverse('rapport_list'))

		self.assertContains(
			response,
			'formaction="/rapports/excel/"',
		)
		self.assertContains(
			response,
			'formaction="/rapports/pdf/"',
		)

	def test_consultation_role_has_no_pdf_footer_link(self):
		user = User.objects.create_user(
			username='consultation_no_pdf',
			password='password123',
			role='consultation',
		)
		self.client.force_login(user)

		response = self.client.get(reverse('rapport_list'))

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'href="/rapports/excel/"')
		self.assertNotContains(response, 'href="/rapports/pdf/"')

	def test_invalid_monthly_and_yearly_filters_are_rejected_on_all_report_routes(self):
		user = User.objects.create_user(
			username='admin_invalid_filters',
			password='password123',
			role='admin',
		)
		self.client.force_login(user)
		invalid_filters = [
			{'type_rapport': 'mensuel', 'mois': '', 'annee': '2026'},
			{'type_rapport': 'mensuel', 'mois': '13', 'annee': '2026'},
			{'type_rapport': 'annuel', 'annee': ''},
			{'type_rapport': 'annuel', 'annee': 'not-a-year'},
		]

		for route in ('rapport_list', 'rapport_excel', 'rapport_pdf'):
			for filters in invalid_filters:
				with self.subTest(route=route, filters=filters):
					response = self.client.get(reverse(route), filters)
					self.assertEqual(response.status_code, 400)


class RapportFiltresTests(TestCase):

	def setUp(self):
		self.site = Site.objects.create(
			nom='Site graphique',
			commune='Commune graphique',
			region='Region graphique',
		)
		self.espece = Espece.objects.create(nom='Espece graphique')
		Debarquement.objects.bulk_create([
			Debarquement(
				date='2026-01-05',
				site=self.site,
				espece=self.espece,
				quantite=100,
			),
			Debarquement(
				date='2026-01-05',
				site=self.site,
				espece=self.espece,
				quantite=50,
			),
			Debarquement(
				date='2026-01-06',
				site=self.site,
				espece=self.espece,
				quantite=75,
			),
			Debarquement(
				date='2026-02-01',
				site=self.site,
				espece=self.espece,
				quantite=200,
			),
		])

	def test_evolution_journaliere_regroupe_par_date(self):
		self.assertEqual(
			construire_evolution(
				Debarquement.objects.all(),
				'journalier',
			),
			[
				{'label': '05/01/2026', 'total': 150},
				{'label': '06/01/2026', 'total': 75},
				{'label': '01/02/2026', 'total': 200},
			],
		)

	def test_evolution_mensuelle_regroupe_par_jour(self):
		self.assertEqual(
			construire_evolution(
				Debarquement.objects.filter(date__month=1),
				'mensuel',
			),
			[
				{'label': '05/01', 'total': 150},
				{'label': '06/01', 'total': 75},
			],
		)

	def test_evolution_annuelle_regroupe_par_mois(self):
		self.assertEqual(
			construire_evolution(
				Debarquement.objects.all(),
				'annuel',
			),
			[
				{'label': '01/2026', 'total': 225},
				{'label': '02/2026', 'total': 200},
			],
		)

	def test_evolution_respecte_les_filtres_site_et_espece(self):
		autre_site = Site.objects.create(
			nom='Autre site graphique',
			commune='Autre commune graphique',
			region='Autre region graphique',
		)
		autre_espece = Espece.objects.create(nom='Autre espece graphique')
		Debarquement.objects.create(
			date='2026-01-05',
			site=autre_site,
			espece=self.espece,
			quantite=900,
		)
		Debarquement.objects.create(
			date='2026-01-05',
			site=self.site,
			espece=autre_espece,
			quantite=800,
		)

		request = RequestFactory().get(
			'/rapports/',
			{
				'type_rapport': 'annuel',
				'annee': '2026',
				'site': self.site.pk,
				'espece': self.espece.pk,
			},
		)
		debarquements, type_rapport, *_ = appliquer_filtres(request)

		self.assertEqual(
			construire_evolution(debarquements, type_rapport),
			[
				{'label': '01/2026', 'total': 225},
				{'label': '02/2026', 'total': 200},
			],
		)

	def test_exports_respectent_le_filtre_espece(self):
		site = Site.objects.create(
			nom='Site export',
			commune='Commune export',
			region='Region export',
		)
		espece_cible = Espece.objects.create(
			nom='Espece exportee'
		)
		autre_espece = Espece.objects.create(
			nom='Espece exclue'
		)
		Debarquement.objects.create(
			date='2026-03-01',
			site=site,
			espece=espece_cible,
			quantite=18,
		)
		Debarquement.objects.create(
			date='2026-03-02',
			site=site,
			espece=autre_espece,
			quantite=40,
		)

		request = RequestFactory().get(
			'/rapports/excel/',
			{'espece': espece_cible.pk},
		)
		request.user = type(
			'ExportUser',
			(),
			{'is_authenticated': True, 'role': 'admin'},
		)()

		excel_response = rapport_excel(request)
		workbook = load_workbook(BytesIO(excel_response.content))
		rows = list(workbook.active.iter_rows(values_only=True))
		values = [value for row in rows for value in row if value is not None]

		self.assertIn('Espece exportee', values)
		self.assertIn(18, values)
		self.assertNotIn('Espece exclue', values)
		self.assertNotIn(40, values)

		request.path = '/rapports/pdf/'
		pdf_response = rapport_pdf(request)

		self.assertEqual(pdf_response.status_code, 200)
		self.assertEqual(pdf_response['Content-Type'], 'application/pdf')
		self.assertIn(
			'filename="rapport_activites.pdf"',
			pdf_response['Content-Disposition'],
		)
		self.assertTrue(pdf_response.content.startswith(b'%PDF-'))
		self.assertGreater(len(pdf_response.content), 10000)

	def test_export_excel_contient_les_sections_du_rapport_d_activites(self):
		request = RequestFactory().get(
			'/rapports/excel/',
			{
				'type_rapport': 'annuel',
				'annee': '2026',
				'site': self.site.pk,
				'espece': self.espece.pk,
			},
		)
		request.user = type(
			'ExportUser',
			(),
			{'is_authenticated': True, 'role': 'admin'},
		)()

		response = rapport_excel(request)
		workbook = load_workbook(BytesIO(response.content))
		worksheet = workbook.active
		values = [
			value
			for row in worksheet.iter_rows(values_only=True)
			for value in row
			if value is not None
		]

		self.assertEqual(
			worksheet['A9'].value,
			'RAPPORT D’ACTIVITÉS ANNUEL — 2026',
		)
		self.assertIn('1. PERSONNEL', values)
		self.assertIn('6. SITUATION DES CARTES PROFESSIONNELLES ET PERMIS DE PÊCHE', values)
		self.assertIn('12. PRODUCTION DE LA PÊCHE LOCALE — MISES À TERRE', values)
		self.assertIn('18. CONTRAINTES LIÉES AU FONCTIONNEMENT DU SERVICE', values)
		self.assertIn('Espece graphique', values)
		self.assertIn(425, values)
		self.assertGreater(len(worksheet._images), 0)
		self.assertEqual(
			response['Content-Disposition'],
			'attachment; filename="rapport_activites.xlsx"',
		)

	def test_export_mensuel_conserve_le_mois_selectionne(self):
		request = RequestFactory().get(
			'/rapports/excel/',
			{
				'type_rapport': 'mensuel',
				'mois': '1',
				'annee': '2026',
			},
		)
		request.user = type(
			'ExportUser',
			(),
			{'is_authenticated': True, 'role': 'admin'},
		)()

		response = rapport_excel(request)
		worksheet = load_workbook(BytesIO(response.content)).active
		values = [
			value
			for row in worksheet.iter_rows(values_only=True)
			for value in row
			if value is not None
		]

		self.assertIn(
			'RAPPORT D’ACTIVITÉS MENSUEL — JANVIER 2026',
			values,
		)
		self.assertIn(225, values)
		self.assertNotIn(200, values)

	def test_filtre_espece_change_les_donnees_et_le_total(self):
		site = Site.objects.create(
			nom='Site test',
			commune='Commune test',
			region='Region test',
		)
		espece_cible = Espece.objects.create(
			nom='Espece selectionnee'
		)
		autre_espece = Espece.objects.create(
			nom='Autre espece',
		)

		debarquement_cible = Debarquement.objects.create(
			date='2026-02-10',
			site=site,
			espece=espece_cible,
			quantite=12,
		)
		Debarquement.objects.create(
			date='2026-02-11',
			site=site,
			espece=autre_espece,
			quantite=30,
		)

		request = RequestFactory().get(
			'/rapports/',
			{
				'espece': espece_cible.pk,
			}
		)

		debarquements, *_ = appliquer_filtres(request)

		self.assertEqual(list(debarquements), [debarquement_cible])
		self.assertEqual(
			debarquements.aggregate(total=Sum('quantite'))['total'],
			12,
		)

	def test_filtre_site_et_espece(self):
		site_cible = Site.objects.create(
			nom='Site cible',
			commune='Commune cible',
			region='Region cible',
		)
		autre_site = Site.objects.create(
			nom='Autre site',
			commune='Autre commune',
			region='Autre region',
		)
		espece_cible = Espece.objects.create(
			nom='Espece cible'
		)
		autre_espece = Espece.objects.create(
			nom='Autre espece'
		)

		debarquement_cible = Debarquement.objects.create(
			date='2026-01-15',
			site=site_cible,
			espece=espece_cible,
			quantite=10,
		)
		Debarquement.objects.create(
			date='2026-01-15',
			site=autre_site,
			espece=espece_cible,
			quantite=20,
		)
		Debarquement.objects.create(
			date='2026-01-15',
			site=site_cible,
			espece=autre_espece,
			quantite=30,
		)

		request = RequestFactory().get(
			'/rapports/',
			{
				'type_rapport': 'annuel',
				'annee': '2026',
				'site': site_cible.pk,
				'espece': espece_cible.pk,
			}
		)

		debarquements, *_ = appliquer_filtres(request)

		self.assertEqual(
			list(debarquements),
			[debarquement_cible],
		)
